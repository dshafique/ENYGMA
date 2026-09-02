"""Documents ENYGMA hands him.

He asked for a markdown file and got a fenced code block on a phone screen. The
fix is only worth anything if the file opens, so every renderer here is checked
by reading its output back with the library that reads that format. A file that
downloads and will not open is worse than no file: it looks like it worked.
"""
import os, sys, pathlib, tempfile, io, subprocess

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
os.environ.setdefault("ENYGMA_SESSION_SECRET", "test-secret")
os.environ.setdefault("ENYGMA_INSECURE_COOKIES", "1")

from src import db, config as cfg                # noqa: E402
from src.ingest import upload                    # noqa: E402

DOC = {
    "title": "Lucelabs Hardware Logbook",
    "subtitle": "Tent config v1.0",
    "blocks": [
        {"type": "heading", "level": 1, "text": "Current deployment"},
        {"type": "paragraph", "text": "The ESP32 is on the tent rail, over USB-C."},
        {"type": "bullets", "items": ["EZO HUM at 0x27", "EZO ENV at 0x1A"]},
        {"type": "numbered", "items": ["Calibrate", "Log", "Commit"]},
        {"type": "table", "columns": ["Sensor", "Address", "Owner"],
         "rows": [["EZO HUM", "0x27", "Joseph"], ["EZO CO2", "0x61", "Jacob"]]},
        {"type": "code", "text": "i2cdetect -y 1"},
        {"type": "quote", "text": "You need to know what the working state was."},
    ],
}


def setup_module(_):
    tmp = pathlib.Path(tempfile.mkdtemp())
    cfg.DB_PATH = tmp / "t.db"; cfg.DATA_DIR = tmp
    db.DB_PATH = cfg.DB_PATH;  db.DATA_DIR = cfg.DATA_DIR
    upload.UPLOADS = tmp / "uploads"; upload.BASE_DIR = tmp
    db.migrate()
    from src import made
    made.BASE_DIR = tmp
    made.MADE = tmp / "made"


def doc():
    from src import documents
    return documents.normalise(DOC)


# --------------------------------------------------------- what he asked for

def test_the_format_comes_from_how_he_asked():
    from src import documents
    cases = {
        "Can you make me a .md file": "md",
        "make me a markdown file for this": "md",
        "put that in a spreadsheet": "xlsx",
        "can I get this as an excel file": "xlsx",
        "turn this into a slide deck": "pptx",
        "make a powerpoint out of it": "pptx",
        "give me a word doc": "docx",
        "export this as a PDF": "pdf",
        "save that as an html page": "html",
        "make me a file with all of that": "md",
        "write up a report on the tent build": "md",
    }
    for text, expected in cases.items():
        assert documents.detect(text) == expected, f"{text!r} -> {documents.detect(text)}"


def test_ordinary_questions_do_not_produce_files():
    """The worst failure here is answering a question with a download."""
    from src import documents
    for text in ["What is a retained message?",
                 "How do I keep track of the work I do?",
                 "Why did the export break?",
                 "What did we decide in the meeting on Tuesday?",
                 "Tell me about the caching setup"]:
        assert documents.detect(text) is None, f"{text!r} wanted a file"


# ------------------------------------------------------- the files then open

def test_the_word_document_opens_and_has_the_table_in_it():
    from docx import Document
    from src import documents
    d = Document(io.BytesIO(documents.render(doc(), "docx")))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Lucelabs Hardware Logbook" in text
    assert "The ESP32 is on the tent rail, over USB-C." in text
    assert len(d.tables) == 1
    assert [c.text for c in d.tables[0].rows[0].cells] == ["Sensor", "Address", "Owner"]
    assert d.tables[0].rows[1].cells[0].text == "EZO HUM"


def test_the_deck_opens_and_a_heading_starts_a_new_slide():
    from pptx import Presentation
    from src import documents
    p = Presentation(io.BytesIO(documents.render(doc(), "pptx")))
    titles = [s.shapes.title.text for s in p.slides if s.shapes.title]
    assert titles == ["Lucelabs Hardware Logbook", "Current deployment"]
    said = "\n".join(sh.text_frame.text for s in p.slides
                     for sh in s.shapes if sh.has_text_frame)
    assert "EZO HUM at 0x27" in said


def test_the_spreadsheet_opens_and_a_table_gets_its_own_sheet():
    """The only reason to ask for a spreadsheet is the rows."""
    from openpyxl import load_workbook
    from src import documents
    wb = load_workbook(io.BytesIO(documents.render(doc(), "xlsx")))
    assert "Document" in wb.sheetnames and len(wb.sheetnames) == 2
    ws = wb[wb.sheetnames[1]]
    assert [c.value for c in ws[1]] == ["Sensor", "Address", "Owner"]
    assert [c.value for c in ws[2]] == ["EZO HUM", "0x27", "Joseph"]
    assert ws.freeze_panes == "A2"


def test_the_pdf_is_a_real_pdf_with_selectable_text():
    """Not a picture of a document. He has to be able to search it."""
    from src import documents
    raw = documents.render(doc(), "pdf")
    assert raw.startswith(b"%PDF-")
    tmp = pathlib.Path(tempfile.mkdtemp()) / "t.pdf"
    tmp.write_bytes(raw)
    out = subprocess.run(["pdftotext", str(tmp), "-"], capture_output=True, text=True)
    assert out.returncode == 0
    assert "Lucelabs" in out.stdout and "EZO HUM" in out.stdout


def test_the_web_page_needs_nothing_from_the_internet():
    """It has to open from a Downloads folder on a phone with no signal."""
    from src import documents
    html = documents.render(doc(), "html").decode()
    assert "http://" not in html and "https://" not in html
    assert "<table" in html and "Lucelabs Hardware Logbook" in html


def test_the_markdown_is_markdown():
    from src import documents
    md = documents.render(doc(), "md").decode()
    assert md.startswith("# Lucelabs Hardware Logbook")
    assert "| Sensor | Address | Owner |" in md
    assert "- EZO HUM at 0x27" in md
    assert "1. Calibrate" in md
    assert "```" in md


def test_html_escapes_rather_than_executes():
    from src import documents
    hostile = documents.normalise({
        "title": "<script>alert(1)</script>",
        "blocks": [{"type": "paragraph", "text": "<img src=x onerror=alert(1)>"}]})
    html = documents.render(hostile, "html").decode()
    # The test is whether the browser would run it, not whether the characters
    # appear: escaped, "onerror=alert(1)" is just words on a page.
    assert "<script>alert(1)</script>" not in html
    assert "<img src=x" not in html
    assert "&lt;script&gt;" in html and "&lt;img src=x" in html


# ------------------------------------------------------------- bad material

def test_a_malformed_structure_renders_instead_of_raising():
    """The model will eventually return something the wrong shape. Every format
    still has to produce a file rather than a traceback."""
    from src import documents
    messy = documents.normalise({
        "title": None,
        "blocks": [
            {"type": "nonsense", "text": "still a paragraph"},
            {"type": "bullets"},                       # no items
            {"type": "table", "rows": []},             # no rows
            {"type": "heading"},                       # no text
            "not even a dict",
            {"type": "paragraph", "text": "  "},       # empty
            {"type": "table", "rows": [["a", "b"]]},   # rows, no columns
        ]})
    assert messy["title"] == "Untitled"
    for fmt in documents.FORMATS:
        assert len(documents.render(messy, fmt)) > 0, fmt


def test_an_empty_document_is_refused_not_shipped():
    """A file that opens on nothing looks like the feature worked."""
    import pytest
    from src import documents
    cfg.config.PIPELINE = "gemini"

    class Empty:
        model = "fake"
        def _ask(self, parts, schema=None): return '{"title": "x", "blocks": []}'
    try:
        with pytest.raises(ValueError):
            documents.compose("make me a doc", backend=Empty())
    finally:
        cfg.config.PIPELINE = "stub"


# ------------------------------------------------------------ kept, and found

def test_the_file_is_kept_in_the_thread_and_in_the_library():
    from src import chat, made, library
    thread_id = chat.start("Tent build")
    row = made.create("Make me a .md file for the tent logbook", "md",
                      context="The ESP32 is on the tent rail.", thread_id=thread_id)
    assert row["filename"].endswith(".md")
    assert [f["id"] for f in made.for_thread(thread_id)] == [row["id"]]

    path, meta = made.blob(row["id"])
    assert path.exists() and path.read_bytes()[:1] == b"#"
    assert meta["mime"] == "text/markdown"

    assert row["library_document_id"], "it never reached the Library"
    assert library.get(row["library_document_id"]) is not None


def test_asking_for_a_file_in_chat_returns_one():
    """The whole point: he types the sentence, he gets the file."""
    from src import chat, made
    thread_id = chat.start("New conversation")
    result = chat.say(thread_id, "Can you make me a .md file for the tent build")
    assert result.get("made"), "no file came back"
    assert result["made"]["format"] == "md"
    assert ".md" in result["reply"]
    assert made.for_thread(thread_id)


def test_asking_a_question_in_chat_returns_an_answer_not_a_file():
    from src import chat
    thread_id = chat.start("New conversation")
    result = chat.say(thread_id, "How should I keep track of what I do?")
    assert result.get("made") is None


def test_the_same_document_twice_is_one_file_on_disk():
    from src import documents, made
    data = documents.render(doc(), "md")
    first, _ = made._store(data, "md")
    second, _ = made._store(data, "md")
    assert first == second


def test_the_file_is_named_for_the_subject_not_for_the_sentence():
    """"Can you make me a .md file for the tent logbook" is a request about the
    tent logbook. Naming the file after the whole sentence gives him a Downloads
    folder full of can-you-make-me-a-md-file-for.md."""
    from src import documents
    cases = {
        "Can you make me a .md file for the Lucelabs tent logbook": "Lucelabs tent logbook",
        "make a powerpoint about the sensor calibration": "sensor calibration",
        "give me a word doc of the weekly numbers": "weekly numbers",
        # The format is on the back of these, not the front.
        "turn the tent logbook into a slide deck": "tent logbook",
        "export the tent logbook as a pdf": "tent logbook",
        "save the tent logbook as an html page": "tent logbook",
        # "build" is a word he asks with AND the name of the thing. Stripping it
        # from both ends turned this into "tent".
        "write up a report on the tent build": "tent build",
        # No subject in the sentence at all: the material is the subject.
        "put that in a spreadsheet": "Document",
        "Turn this answer into a document": "Document",
    }
    for asked, expected in cases.items():
        assert documents.title_from_request(asked) == expected, asked


def test_pressing_send_does_not_move_the_send_button():
    """The bug this guards: hiding the tab bar on focus and restoring it on blur
    meant pressing Send blurred the box, which put the bar back, which moved the
    composer sixty pixels between his finger going down and coming up. The
    button slid out from under his thumb as he pressed it and nothing happened
    at all."""
    js = (pathlib.Path(__file__).resolve().parent.parent
          / "src/static/js/app.js").read_text()
    chat = js[js.index("-------- chat */"):js.index("------- settings */")]
    # Pressing Send must not take focus off the box.
    assert 'button?.addEventListener("pointerdown", (e) => e.preventDefault());' in chat
    # And the bar comes back only once focus has genuinely left the composer.
    assert '"focusout"' in chat and '"focusin"' in chat
    assert "composer.contains(document.activeElement)" in chat
    assert 'body?.addEventListener("blur"' not in chat, "back to blurring on the box"


def test_nothing_is_stored_outside_a_path_systemd_makes_writable():
    """The bug this exists for: made documents were stored in app/made, a new
    directory beside data/ rather than inside it. The unit runs with
    ProtectSystem=strict and ProtectHome=read-only and grants exactly two
    writable paths, so that directory was read-only at the kernel level whatever
    its permissions said. Every test passed. It failed on the Spark the first
    time he asked for a file, with Errno 30.

    No runtime test could catch it: the suite runs in a temp directory with no
    sandbox, and it moves these very paths to get there. So this reads the
    source instead and asks a structural question -- does any module hang a
    store off BASE_DIR that the unit has not made writable.
    """
    import re
    root = pathlib.Path(__file__).resolve().parent.parent

    granted = []
    for line in (root / "deploy/enygma.service").read_text().splitlines():
        if line.startswith("ReadWritePaths="):
            granted += line.split("=", 1)[1].split()
    assert granted, "the unit grants no writable paths at all"
    # What the granted paths are called, relative to the app directory.
    writable = {pathlib.PurePosixPath(g).name for g in granted}
    assert {"data", "uploads"} <= writable, granted

    offenders = []
    for path in sorted((root / "src").rglob("*.py")):
        for n, line in enumerate(path.read_text().splitlines(), 1):
            # A store declared straight off the application root.
            found = re.search(r'^\s*\w+\s*=\s*BASE_DIR\s*/\s*"([^"]+)"', line)
            if found and found.group(1) not in writable:
                offenders.append(f"{path.relative_to(root)}:{n}  {line.strip()}")
    assert not offenders, (
        "these are written to but are read-only under the service unit:\n  "
        + "\n  ".join(offenders)
        + f"\nWritable: {sorted(writable)}. Put it under data/.")

    # And the documents store specifically lives under data/, in the source,
    # not merely wherever a test happened to point it.
    made_src = (root / "src/made.py").read_text()
    assert 'MADE = DATA_DIR / "made"' in made_src, \
        "the made-documents store moved back out of data/"


def test_a_document_survives_the_data_directory_moving_to_another_disk():
    """Paths were stored relative to the application root, so the moment data/
    was not inside app/ -- which is what you do when the disk fills -- every
    document failed to save with a path error that reads like a renderer bug."""
    import importlib
    from src import made, documents
    tmp = pathlib.Path(tempfile.mkdtemp())
    elsewhere = tmp / "on-another-disk"
    keep = (made.BASE_DIR, made.DATA_DIR, made.MADE)
    try:
        made.BASE_DIR = tmp / "app"
        made.DATA_DIR = elsewhere
        made.MADE = elsewhere / "made"
        stored, _ = made._store(documents.render(doc(), "md"), "md")
        assert not stored.startswith("/"), stored
        assert (made.DATA_DIR / stored).exists(), stored
    finally:
        made.BASE_DIR, made.DATA_DIR, made.MADE = keep


def test_documents_written_by_an_older_build_still_download():
    """Those rows hold "data/made/...", relative to the application root."""
    from src import made, documents
    tmp = pathlib.Path(tempfile.mkdtemp())
    keep = (made.BASE_DIR, made.DATA_DIR, made.MADE)
    try:
        made.BASE_DIR = tmp
        made.DATA_DIR = tmp / "data"
        made.MADE = made.DATA_DIR / "made"
        made.MADE.mkdir(parents=True)
        old = made.MADE / "deadbeef.md"
        old.write_bytes(b"# from an older build\n")
        with db.cursor() as conn:
            conn.execute(
                "INSERT INTO made_documents (format, style, title, filename, mime, "
                "path, bytes) VALUES ('md','plain','Old','old.md','text/markdown',"
                "'data/made/deadbeef.md', 22)")
            made_id = conn.execute("SELECT last_insert_rowid() AS i").fetchone()["i"]
        found = made.blob(made_id)
        assert found is not None, "an older build's document became undownloadable"
        assert found[0].read_bytes().startswith(b"# from an older build")
    finally:
        made.BASE_DIR, made.DATA_DIR, made.MADE = keep


def test_a_request_with_no_subject_is_named_for_the_thread():
    """"put that in a spreadsheet" names nothing. A folder of document.xlsx and
    document.docx is a folder he cannot search."""
    from src import chat, made
    thread_id = chat.start("New conversation")
    chat.say(thread_id, "How should I track the Lucelabs tent build?")
    chat.say(thread_id, "put that in a spreadsheet")
    files = made.for_thread(thread_id)
    assert files, "no file was made"
    assert not files[-1]["filename"].startswith("document."), files[-1]["filename"]
    assert "tent" in files[-1]["filename"].lower(), files[-1]["filename"]


def test_a_borrowed_thread_title_is_trimmed_like_any_other_request():
    """A thread is named after its first message, and that message is often
    itself a request for a file. Borrowing it raw hands back exactly the
    sentence the trimming existed to remove."""
    from src import documents
    got = documents._stub("put that in a spreadsheet", "",
                          "Make me a .md file for the tent logbook")
    assert got["title"] == "tent logbook", got["title"]


def test_an_unnamed_thread_is_not_a_title():
    from src import documents
    assert documents._stub("turn this into a deck", "", "New conversation")["title"] \
        == "Document"
